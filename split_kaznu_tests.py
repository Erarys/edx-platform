from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


DOWNLOADS = Path.home() / "Downloads"
SOURCE_FILE = DOWNLOADS / "KazNU_ST2026_questions.xlsx"
TEMPLATE_FILE = DOWNLOADS / "questions-template.xlsx"
OUTPUT_DIR = DOWNLOADS / "Gaj_tests"
SHEET_NAME = "Questions"

# Диапазоны — включительные, с нумерацией вопросов от 0.
# Тест 1.4 отсутствует и поэтому намеренно не создаётся.
# Вопрос 446 входит и в 6.1, и в итоговый тест.
TESTS = [
    ("1.1", 0, 31),
    ("1.2", 32, 59),
    ("1.3", 60, 88),
    ("2.1", 90, 116),
    ("2.2", 117, 146),
    ("3.1", 147, 176),
    ("3.2", 177, 206),
    ("3.3", 207, 236),
    ("3.4", 237, 266),
    ("4.1", 267, 296),
    ("4.2", 297, 326),
    ("5.1", 327, 356),
    ("5.2", 357, 386),
    ("5.3", 387, 416),
    ("6.1", 417, 446),
    ("ҚОРЫТЫНДЫ ТЕСТ", 446, None),
]


def read_questions():
    workbook = load_workbook(SOURCE_FILE, data_only=False, read_only=True)
    worksheet = workbook[SHEET_NAME]
    rows = list(
        worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
            min_col=1,
            max_col=8,
            values_only=True,
        )
    )
    workbook.close()

    if len(rows) <= 446:
        raise ValueError(f"В исходнике найдено только {len(rows)} вопросов")
    return rows


def create_test_file(code, selected_rows):
    workbook = load_workbook(TEMPLATE_FILE)
    worksheet = workbook[SHEET_NAME]

    # Запоминаем оформление первой строки данных шаблона.
    sample_styles = [copy(worksheet.cell(2, column)._style) for column in range(1, 9)]
    sample_height = worksheet.row_dimensions[2].height

    # Удаляем демонстрационные вопросы, сохраняя заголовок и структуру шаблона.
    if worksheet.max_row > 1:
        worksheet.delete_rows(2, worksheet.max_row - 1)

    for row_number, row_values in enumerate(selected_rows, start=2):
        for column, value in enumerate(row_values, start=1):
            cell = worksheet.cell(row=row_number, column=column, value=value)
            cell._style = copy(sample_styles[column - 1])
        if sample_height is not None:
            worksheet.row_dimensions[row_number].height = sample_height

    last_row = len(selected_rows) + 1
    table_ref = f"A1:H{last_row}"

    tables = list(worksheet.tables.values())
    if tables:
        tables[0].ref = table_ref
    else:
        table = Table(displayName="QuestionsTable", ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    output_file = OUTPUT_DIR / f"Gaj{code}.xlsx"
    workbook.save(output_file)
    workbook.close()
    return output_file


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    questions = read_questions()

    created = []
    for code, start, end in TESTS:
        inclusive_end = len(questions) - 1 if end is None else end
        selected = questions[start : inclusive_end + 1]
        expected = inclusive_end - start + 1
        if len(selected) != expected:
            raise ValueError(
                f"{code}: ожидалось {expected} вопросов, получено {len(selected)}"
            )
        created.append(create_test_file(code, selected))

    print(f"Готово: создано {len(created)} файлов в {OUTPUT_DIR}")
    for file in created:
        print(file.name)


if __name__ == "__main__":
    main()
